"""Synthetic support-ticket log with exact ground truth.

The assignment describes the input as a "ticket log", and the tool must work
beyond the Indian prospectus. This builds a .docx of support tickets in US, UK
and Indian formats. Every PII value is inserted by this script, so the gold
labels are exact. Non-PII look-alikes (ticket/order numbers, versions,
ordinary dates, amounts, acronyms) are mixed in to measure precision.

The set was generated and frozen *before* any change made for it.

Usage: python eval/tickets/generate.py   (writes eval/tickets/tickets.docx,
       sample.jsonl and ground_truth.jsonl)
"""
from __future__ import annotations

import json
import random
import re
import sys

import docx
from faker import Faker

sys.path.insert(0, ".")
from redactor.docx_io import DocxDocument  # noqa: E402

OUT = "eval/tickets/"
SEED = 2026
LOCALES = ["en_US"] * 6 + ["en_GB"] * 2 + ["en_IN"] * 2
PRODUCTS = ["VPN client", "billing portal", "mobile app", "SSO login", "API gateway", "CRM sync", "invoice export"]
ISSUES = [
    "cannot log in after the 2FA reset", "gets HTTP 502 on checkout", "sees a duplicate charge",
    "reports the export times out after 30 seconds", "wants the SLA breach credit applied",
    "asks why the PDF invoice shows the wrong VAT", "lost access after the password policy change",
]


class Line:
    """Builds one paragraph and remembers where each PII value was put."""

    def __init__(self):
        self.text, self.spans = "", []

    def add(self, s: str, label: str | None = None) -> "Line":
        if label:
            self.spans.append((label, s))
        self.text += s
        return self


def phone(f: Faker) -> str:
    return re.sub(r"x\d+$", "", f.phone_number())  # drop extensions ("x9150")


def card(r: random.Random, f: Faker) -> str:
    n = f.credit_card_number(card_type=r.choice(["visa16", "mastercard", "amex", "discover"]))
    if r.random() < 0.5 and len(n) == 16:
        n = " ".join(n[i:i + 4] for i in range(0, 16, 4))
    return n


def ticket(r: random.Random, i: int) -> list[Line]:
    loc = r.choice(LOCALES)
    f = Faker(loc)
    f.seed_instance(r.randrange(10**9))
    first = f.first_name()
    name = f"{first} {f.last_name()}"  # no "Dr."/"MD" noise in the gold span
    email = re.sub(r"[^a-z.]", "", name.lower().replace(" ", ".")) + "@" + f.free_email_domain()
    company = f.company()
    tid, order = f"TCK-{100000 + i}", f"ORD-{r.randrange(10**7):07d}"
    created = f.date_between("-1y", "today").strftime("%B %d, %Y")
    lines = [
        Line().add(f"Ticket {tid} | Priority P{r.randint(1, 4)} | Opened {created} | Product: {r.choice(PRODUCTS)} v{r.randint(1, 9)}.{r.randint(0, 20)}.{r.randint(0, 99)}"),
        Line().add("Customer: ").add(name, "PERSON").add(" (").add(email, "EMAIL").add("), ").add(company, "ORGANIZATION"),
        Line().add(f"{first}", "PERSON").add(f" {r.choice(ISSUES)}. Related order {order}, amount ${r.randint(20, 900)}.{r.randint(10, 99)}."),
    ]
    kind = r.random()
    if kind < 0.35:
        lines.append(Line().add("Callback number: ").add(phone(f), "PHONE").add(f". Best time after {r.randint(1, 5)} pm."))
    elif kind < 0.6:
        addr = ", ".join(f.address().splitlines())
        lines.append(Line().add("Shipping address on file: ").add(addr, "ADDRESS").add("."))
    elif kind < 0.75:
        lines.append(Line().add("Login attempts came from IP ").add(f.ipv4_public(), "IP_ADDRESS").add(f" at {r.randint(0, 23):02d}:{r.randint(0, 59):02d} UTC."))
    elif kind < 0.85:
        lines.append(Line().add("Identity check passed: date of birth ").add(f.date_of_birth(minimum_age=18, maximum_age=80).strftime("%m/%d/%Y"), "DATE_OF_BIRTH").add(" confirmed."))
    elif kind < 0.93 and loc == "en_US":
        lines.append(Line().add("Customer pasted their SSN ").add(f.ssn(), "SSN").add(" into the chat; please purge it."))
    else:
        lines.append(Line().add("Refund to card ").add(card(r, f), "CREDIT_CARD").add(f" approved by agent #{r.randint(100, 999)}."))
    if r.random() < 0.3:
        agent = Faker("en_US")
        agent.seed_instance(r.randrange(10**9))
        lines.append(Line().add("Escalated to ").add(f"{agent.first_name()} {agent.last_name()}", "PERSON").add(" (Tier 2). See KB article 4471 and RFC 7231."))
    lines.append(Line().add(f"Status: {r.choice(['Resolved', 'Pending customer', 'Escalated', 'Closed'])}."))
    return lines


def main() -> None:
    r = random.Random(SEED)
    d = docx.Document()
    d.add_paragraph("Support ticket export - Q3 review. Internal use only.")
    all_lines = []
    for i in range(120):
        for line in ticket(r, i):
            d.add_paragraph(line.text)
            all_lines.append(line)
    path = OUT + "tickets.docx"
    d.save(path)
    # uids come from the same reader the tool uses
    paras = DocxDocument(path).paragraphs[1:]
    assert [p.text for p in paras] == [l.text for l in all_lines]
    with open(OUT + "sample.jsonl", "w") as s, open(OUT + "ground_truth.jsonl", "w") as g:
        for p, line in zip(paras, all_lines):
            s.write(json.dumps({"paragraph": p.uid, "stratum": "all", "text": p.text}, ensure_ascii=False) + "\n")
            for label, value in line.spans:
                g.write(json.dumps({"paragraph": p.uid, "label": label, "text": value}, ensure_ascii=False) + "\n")
    print(f"{len(all_lines)} paragraphs, {sum(len(l.spans) for l in all_lines)} PII spans -> {path}")


if __name__ == "__main__":
    main()
