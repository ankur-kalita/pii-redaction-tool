import docx

from redactor.docx_io import DocxDocument, Replacement, apply_replacements


def _make_doc(tmp_path):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Contact Sarth")
    p.add_run("ak Malva").bold = True
    p.add_run("dkar today")
    table = d.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Email: a@b.com"
    d.sections[0].footer.paragraphs[0].text = "Footer Rohan Dey"
    path = tmp_path / "in.docx"
    d.save(path)
    return path


def test_reads_body_tables_and_footers(tmp_path):
    doc = DocxDocument(str(_make_doc(tmp_path)))
    texts = [p.text for p in doc.paragraphs]
    assert "Contact Sarthak Malvadkar today" in texts
    assert "Email: a@b.com" in texts
    assert "Footer Rohan Dey" in texts


def test_replacement_across_split_runs_keeps_formatting(tmp_path):
    src = _make_doc(tmp_path)
    doc = DocxDocument(str(src))
    para = next(p for p in doc.paragraphs if p.text.startswith("Contact"))
    start = para.text.index("Sarthak")
    apply_replacements(para, [Replacement(start, start + len("Sarthak Malvadkar"), "John Doe")])
    assert para.text == "Contact John Doe today"
    out = tmp_path / "out.docx"
    doc.save(str(out))
    runs = docx.Document(out).paragraphs[0].runs
    assert "".join(r.text for r in runs) == "Contact John Doe today"
    assert runs[0].bold is not True  # replacement lives in the first touched run


def test_word_formatting_is_kept_when_word_counts_match(tmp_path):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Client Rashi ")
    p.add_run("Patil").bold = True
    p.add_run(" joined.")
    path = tmp_path / "fmt.docx"
    d.save(path)
    doc = DocxDocument(str(path))
    para = doc.paragraphs[0]
    s = para.text.index("Rashi")
    apply_replacements(para, [Replacement(s, s + len("Rashi Patil"), "Jatin Nanda")])
    out = tmp_path / "fmt_out.docx"
    doc.save(str(out))
    runs = docx.Document(out).paragraphs[0].runs
    assert [(r.text, bool(r.bold)) for r in runs] == [("Client Jatin ", False), ("Nanda", True), (" joined.", False)]


def test_multiple_replacements_in_one_paragraph(tmp_path):
    doc = DocxDocument(str(_make_doc(tmp_path)))
    para = next(p for p in doc.paragraphs if p.text.startswith("Email"))
    s = para.text.index("a@b.com")
    apply_replacements(para, [Replacement(0, 5, "Mail"), Replacement(s, s + 7, "x@y.org")])
    assert para.text == "Mail: x@y.org"
